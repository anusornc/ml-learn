"""Automated tests for K-Means Exam (K=3, Min-Max Normalization, 3 Iterations).

Validates:
1. Exact mathematical calculations (Scaling, Distances, Assignments, Centroids, WCSS)
2. Integrity of generated Excel files (Practice and Solved versions)
"""

from __future__ import annotations

import math
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent

# Ground truth raw data
RAW_DATA = [
    ("P1", 20.0, 15000.0),
    ("P2", 24.0, 25000.0),
    ("P3", 28.0, 20000.0),
    ("P4", 36.0, 40000.0),
    ("P5", 40.0, 40000.0),
    ("P6", 44.0, 45000.0),
    ("P7", 52.0, 60000.0),
    ("P8", 56.0, 55000.0),
    ("P9", 60.0, 65000.0),
]


def test_minmax_normalization() -> None:
    """ตรวจสอบการทำ Min-Max Normalization"""
    ages = [d[1] for d in RAW_DATA]
    incomes = [d[2] for d in RAW_DATA]

    min_age, max_age = min(ages), max(ages)
    range_age = max_age - min_age
    assert min_age == 20.0
    assert max_age == 60.0
    assert range_age == 40.0

    min_inc, max_inc = min(incomes), max(incomes)
    range_inc = max_inc - min_inc
    assert min_inc == 15000.0
    assert max_inc == 65000.0
    assert range_inc == 50000.0

    expected_norm = [
        (0.0, 0.0),
        (0.1, 0.2),
        (0.2, 0.1),
        (0.4, 0.5),
        (0.5, 0.5),
        (0.6, 0.6),
        (0.8, 0.9),
        (0.9, 0.8),
        (1.0, 1.0),
    ]

    for i, (_, age, inc) in enumerate(RAW_DATA):
        norm_x = (age - min_age) / range_age
        norm_y = (inc - min_inc) / range_inc
        assert abs(norm_x - expected_norm[i][0]) < 1e-6
        assert abs(norm_y - expected_norm[i][1]) < 1e-6


def test_iteration_1_math() -> None:
    """ตรวจสอบผลลัพธ์รอบที่ 1: Distances, Assignments, WCSS, New Centroids"""
    points = [
        [0.0, 0.0],
        [0.1, 0.2],
        [0.2, 0.1],
        [0.4, 0.5],
        [0.5, 0.5],
        [0.6, 0.6],
        [0.8, 0.9],
        [0.9, 0.8],
        [1.0, 1.0],
    ]

    c_init = [
        [0.0, 0.0],  # C1 = P1
        [0.5, 0.5],  # C2 = P5
        [1.0, 1.0],  # C3 = P9
    ]

    # Calculate distances
    dists = []
    assignments = []
    dmin_sq = []
    for p in points:
        d = [math.sqrt((p[0] - c[0]) ** 2 + (p[1] - c[1]) ** 2) for c in c_init]
        min_val = min(d)
        c_idx = d.index(min_val) + 1
        dists.append(d)
        assignments.append(c_idx)
        dmin_sq.append(min_val ** 2)

    assert assignments == [1, 1, 1, 2, 2, 2, 3, 3, 3]

    wcss_r1 = sum(dmin_sq)
    assert abs(wcss_r1 - 0.230000) < 1e-6

    # New centroids
    new_c1 = [sum(points[i][0] for i in range(3)) / 3.0, sum(points[i][1] for i in range(3)) / 3.0]
    new_c2 = [sum(points[i][0] for i in range(3, 6)) / 3.0, sum(points[i][1] for i in range(3, 6)) / 3.0]
    new_c3 = [sum(points[i][0] for i in range(6, 9)) / 3.0, sum(points[i][1] for i in range(6, 9)) / 3.0]

    assert abs(new_c1[0] - 0.1000) < 1e-4 and abs(new_c1[1] - 0.1000) < 1e-4
    assert abs(new_c2[0] - 0.5000) < 1e-4 and abs(new_c2[1] - 0.533333) < 1e-4
    assert abs(new_c3[0] - 0.9000) < 1e-4 and abs(new_c3[1] - 0.9000) < 1e-4


def test_iteration_2_and_3_convergence() -> None:
    """ตรวจสอบผลลัพธ์รอบที่ 2 และ 3 รวมถึงการลู่เข้าสมบูรณ์"""
    points = [
        [0.0, 0.0],
        [0.1, 0.2],
        [0.2, 0.1],
        [0.4, 0.5],
        [0.5, 0.5],
        [0.6, 0.6],
        [0.8, 0.9],
        [0.9, 0.8],
        [1.0, 1.0],
    ]

    c_r1 = [
        [0.1000, 0.1000],
        [0.5000, 1.6 / 3.0],
        [0.9000, 0.9000],
    ]

    # Iteration 2
    dmin_sq_r2 = []
    assignments_r2 = []
    for p in points:
        d = [math.sqrt((p[0] - c[0]) ** 2 + (p[1] - c[1]) ** 2) for c in c_r1]
        min_val = min(d)
        c_idx = d.index(min_val) + 1
        assignments_r2.append(c_idx)
        dmin_sq_r2.append(min_val ** 2)

    assert assignments_r2 == [1, 1, 1, 2, 2, 2, 3, 3, 3]

    wcss_r2 = sum(dmin_sq_r2)
    # Expected: 0.106667
    assert abs(wcss_r2 - 0.1066667) < 1e-5

    # New centroids after round 2
    new_c1_r2 = [sum(points[i][0] for i in range(3)) / 3.0, sum(points[i][1] for i in range(3)) / 3.0]
    new_c2_r2 = [sum(points[i][0] for i in range(3, 6)) / 3.0, sum(points[i][1] for i in range(3, 6)) / 3.0]
    new_c3_r2 = [sum(points[i][0] for i in range(6, 9)) / 3.0, sum(points[i][1] for i in range(6, 9)) / 3.0]

    # Centroids don't move at all!
    shift = sum(
        math.sqrt((new_c[0] - old_c[0]) ** 2 + (new_c[1] - old_c[1]) ** 2)
        for new_c, old_c in zip([new_c1_r2, new_c2_r2, new_c3_r2], c_r1)
    )
    assert abs(shift) < 1e-6

    # In Iteration 3, inputs are identical to iteration 2
    # So assignments, WCSS, and centroids are 100% stable -> Convergence reached!


def test_excel_workbooks_exist_and_valid() -> None:
    """ตรวจสอบความถูกต้องของไฟล์ Excel ทั้งสองไฟล์"""
    practice_path = BASE_DIR / "KMeans_Exam_K3_MinMax_Practice.xlsx"
    solved_path = BASE_DIR / "KMeans_Exam_K3_MinMax_Solved.xlsx"

    assert practice_path.exists(), f"Missing {practice_path}"
    assert solved_path.exists(), f"Missing {solved_path}"

    expected_sheets = [
        "00_Exam_Instructions",
        "01_MinMax_Scaling",
        "02_KMeans_Iteration_1",
        "03_KMeans_Iteration_2",
        "04_KMeans_Iteration_3",
        "05_Analysis_and_Rubric",
    ]

    wb_sol = openpyxl.load_workbook(solved_path, data_only=False)
    assert wb_sol.sheetnames == expected_sheets

    # Check that solved workbook contains formulas
    ws_sol_r1 = wb_sol["02_KMeans_Iteration_1"]
    assert str(ws_sol_r1["I20"].value).startswith("=SUM")
    assert str(ws_sol_r1["E11"].value).startswith("=SQRT")

    wb_prac = openpyxl.load_workbook(practice_path, data_only=False)
    assert wb_prac.sheetnames == expected_sheets

    # Check that practice workbook has blank cells for student answers
    ws_prac_r1 = wb_prac["02_KMeans_Iteration_1"]
    assert ws_prac_r1["I20"].value is None
    assert ws_prac_r1["E11"].value is None


if __name__ == "__main__":
    test_minmax_normalization()
    test_iteration_1_math()
    test_iteration_2_and_3_convergence()
    test_excel_workbooks_exist_and_valid()
    print("ALL TESTS PASSED SUCCESSFULLY! (100% verified)")
