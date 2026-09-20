import openpyxl

wb = openpyxl.load_workbook("/Users/anusornchaikaew/Work/lecture/ml-learn/lab-unsupervised-kmeans/KMeans_Exam_K3_MinMax_Solved.xlsx", data_only=False)

print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    # Check for formulas with #REF
    ref_errors = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if isinstance(val, str):
                if "#REF!" in val or "#VALUE!" in val:
                    print(f"  Error at {ws.cell(row=r, column=c).coordinate}: {val}")
                    ref_errors += 1
    if ref_errors == 0:
        print("  No formula syntax errors (#REF!, #VALUE!) found!")

# Let's inspect some key formula cells in 01_MinMax_Scaling
ws_scale = wb['01_MinMax_Scaling']
print("\n01_MinMax_Scaling samples:")
print("  G6 (P1 Age Norm):", ws_scale['G6'].value)
print("  H6 (P1 Inc Norm):", ws_scale['H6'].value)
print("  G14 (P9 Age Norm):", ws_scale['G14'].value)
print("  H14 (P9 Inc Norm):", ws_scale['H14'].value)
print("  F24 (C1 init x):", ws_scale['F24'].value)
print("  G24 (C1 init y):", ws_scale['G24'].value)
print("  F25 (C2 init x):", ws_scale['F25'].value)
print("  G25 (C2 init y):", ws_scale['G25'].value)
print("  F26 (C3 init x):", ws_scale['F26'].value)
print("  G26 (C3 init y):", ws_scale['G26'].value)

# Let's inspect 02_KMeans_Iteration_1
ws_r1 = wb['02_KMeans_Iteration_1']
print("\n02_KMeans_Iteration_1 samples:")
print("  B5 (C1 x):", ws_r1['B5'].value)
print("  C5 (C1 y):", ws_r1['C5'].value)
print("  B6 (C2 x):", ws_r1['B6'].value)
print("  C6 (C2 y):", ws_r1['C6'].value)
print("  B7 (C3 x):", ws_r1['B7'].value)
print("  C7 (C3 y):", ws_r1['C7'].value)
print("  E11 (P1 dist C1):", ws_r1['E11'].value)
print("  F11 (P1 dist C2):", ws_r1['F11'].value)
print("  G11 (P1 dist C3):", ws_r1['G11'].value)
print("  H11 (P1 cluster):", ws_r1['H11'].value)
print("  I11 (P1 dmin sq):", ws_r1['I11'].value)
print("  I20 (WCSS total):", ws_r1['I20'].value)
print("  C24 (New C1 x):", ws_r1['C24'].value)
print("  D24 (New C1 y):", ws_r1['D24'].value)
print("  C25 (New C2 x):", ws_r1['C25'].value)
print("  D25 (New C2 y):", ws_r1['D25'].value)
print("  C26 (New C3 x):", ws_r1['C26'].value)
print("  D26 (New C3 y):", ws_r1['D26'].value)

# Let's inspect 03_KMeans_Iteration_2
ws_r2 = wb['03_KMeans_Iteration_2']
print("\n03_KMeans_Iteration_2 samples:")
print("  B5 (C1 x):", ws_r2['B5'].value)
print("  C5 (C1 y):", ws_r2['C5'].value)
print("  E11 (P1 dist C1):", ws_r2['E11'].value)
print("  H11 (P1 cluster):", ws_r2['H11'].value)
print("  I20 (WCSS total):", ws_r2['I20'].value)
print("  J20 (WCSS diff):", ws_r2['J20'].value)

# Let's inspect 04_KMeans_Iteration_3
ws_r3 = wb['04_KMeans_Iteration_3']
print("\n04_KMeans_Iteration_3 samples:")
print("  B5 (C1 x):", ws_r3['B5'].value)
print("  C5 (C1 y):", ws_r3['C5'].value)
print("  H11 (P1 cluster):", ws_r3['H11'].value)
print("  I20 (WCSS total):", ws_r3['I20'].value)
print("  B30 (Conv reassign):", ws_r3['B30'].value)
print("  B31 (Conv shift):", ws_r3['B31'].value)
print("  B32 (Conv dWCSS):", ws_r3['B32'].value)
