import openpyxl

wb_p = openpyxl.load_workbook("/Users/anusornchaikaew/Work/lecture/ml-learn/lab-unsupervised-kmeans/KMeans_Exam_K3_MinMax_Practice.xlsx", data_only=False)

for sheet_name in wb_p.sheetnames:
    ws = wb_p[sheet_name]
    formula_count = 0
    blank_yellow_count = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                formula_count += 1
            if val is None and cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb == "00FFF3BF":
                blank_yellow_count += 1
    print(f"Sheet {sheet_name}: {formula_count} formulas (should be 0 or only structural), {blank_yellow_count} blank yellow cells for students to answer.")
